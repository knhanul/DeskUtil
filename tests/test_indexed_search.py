import os
import sqlite3
import tempfile
from contextlib import closing
import unittest
from pathlib import Path
from unittest.mock import patch

from doc_search.database.fts5_db import FTS5Database
from doc_search.indexer import DocumentIndexer
from doc_search.search import DocumentSearch


class IndexedSearchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.folder = self.root / 'documents'
        self.folder.mkdir()
        self.db_path = self.root / 'index.db'
        self.db = FTS5Database(str(self.db_path))
        self.db.add_folder(str(self.folder))

    def test_scoped_search_and_stale_file_remain_safe_until_update(self):
        subfolder = self.folder / 'scoped'
        subfolder.mkdir()
        inside = subfolder / 'one.txt'
        outside = self.folder / 'two.txt'
        inside.write_text('sample', encoding='utf-8')
        outside.write_text('sample', encoding='utf-8')
        DocumentIndexer(str(self.db_path)).update_registered({'.txt'})
        self.assertEqual(len(self.db.search_indexed('sample', {'.txt'}, [str(subfolder)])), 1)
        inside.unlink()
        self.assertEqual(len(self.db.search_indexed('sample', {'.txt'}, [str(subfolder)])), 1)
        DocumentIndexer(str(self.db_path)).update_registered({'.txt'})
        self.assertEqual(self.db.search_indexed('sample', {'.txt'}, [str(subfolder)]), [])
        self.assertEqual(len(self.db.search_indexed('sample', {'.txt'}, self.db.list_folders())), 1)

    def test_incremental_update_reuses_unchanged_content_and_prunes_deleted(self):
        source = self.folder / 'policy.txt'
        source.write_text('old policy', encoding='utf-8')
        indexer = DocumentIndexer(str(self.db_path))
        first = indexer.update_registered({'.txt'})
        self.assertEqual(first['new'], 1)
        with patch('doc_search.indexer.get_extractor', side_effect=AssertionError('unchanged file was extracted')):
            unchanged = indexer.update_registered({'.txt'})
        self.assertEqual(unchanged['unchanged'], 1)
        source.write_text('new policy', encoding='utf-8')
        stat = source.stat()
        os.utime(source, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000_000))
        changed = indexer.update_registered({'.txt'})
        self.assertEqual(changed['modified'], 1)
        search = DocumentSearch(str(self.db_path))
        self.assertEqual(search.search_indexed('old', {'.txt'}, self.db.list_folders()), [])
        self.assertEqual(len(search.search_indexed('new', {'.txt'}, self.db.list_folders())), 1)
        source.unlink()
        deleted = indexer.update_registered({'.txt'})
        self.assertEqual(deleted['deleted'], 1)
        self.assertEqual(search.search_indexed('new', {'.txt'}, self.db.list_folders()), [])

    def test_ten_thousand_unchanged_files_do_not_invoke_extractors(self):
        from types import SimpleNamespace
        indexer = DocumentIndexer(str(self.db_path))
        paths = [str(self.folder / f'{index}.txt') for index in range(10_000)]
        snapshot = {path: {'file_size': 12, 'mtime_ns': 99, 'status': 'indexed'} for path in paths}
        indexer.scanner.scan_index_files = lambda extensions, cancelled: (paths, 0, True)
        indexer.db.document_snapshot = lambda folder: snapshot
        indexer.db.save_index_run = lambda unsupported: None
        with patch('doc_search.indexer.os.stat', return_value=SimpleNamespace(st_size=12, st_mtime_ns=99)):
            with patch('doc_search.indexer.get_extractor', side_effect=AssertionError('unchanged file was opened')):
                result = indexer.update_registered({'.txt'})
        self.assertEqual(result['unchanged'], 10_000)
        self.assertEqual(result['new'] + result['modified'] + result['deleted'], 0)

    def test_failed_file_does_not_stop_batch_and_roots_persist(self):
        (self.folder / 'empty.txt').write_text('', encoding='utf-8')
        (self.folder / 'valid.txt').write_text('searchable data', encoding='utf-8')
        indexed = DocumentIndexer(str(self.db_path)).update_registered({'.txt'})
        self.assertEqual((indexed['failed'], indexed['new']), (1, 1))
        self.assertEqual(self.db.list_folders(), FTS5Database(str(self.db_path)).list_folders())
        stats = self.db.index_stats()
        self.assertEqual((stats['indexed'], stats['failed']), (1, 1))
        self.assertTrue(stats['errors'])
        self.assertEqual(len(self.db.search_indexed('searchable', {'.txt'}, self.db.list_folders())), 1)
        self.assertEqual(len(self.db.search_indexed('empty', {'.txt'}, self.db.list_folders(),
                                                   search_content=False, search_filename=True)), 1)

    def test_filename_content_filters_special_characters_and_scope(self):
        source = self.folder / '보험약관_2026.txt'
        source.write_text('보험료(특약) A+B "계약" 10% 상품-설명서', encoding='utf-8')
        self.db.upsert_indexed(str(source), source.read_text(encoding='utf-8'), source.stat())
        search = DocumentSearch(str(self.db_path))
        roots = self.db.list_folders()
        extensions = {'.txt'}
        for term in ('보험료(특약)', 'A+B', '"계약"', '10%', '상품-설명서'):
            with self.subTest(term=term):
                self.assertEqual(len(search.search_indexed(term, extensions, roots)), 1)
        self.assertEqual(len(search.search_indexed('약관', extensions, roots,
                                                   search_content=False, search_filename=True)), 1)
        self.assertEqual(search.search_indexed('약관', extensions, roots,
                                               search_content=True, search_filename=False), [])
        self.assertEqual(search.search_indexed('보험료', extensions, [str(self.root / 'elsewhere')]), [])
        self.assertEqual(search.search_indexed('', extensions, roots), [])
        for term in ('"', '+', '%', '()', '\\', '_', '*', ':', 'a"b', "'", '-'):
            with self.subTest(punctuation=term):
                self.assertIsInstance(search.search_indexed(term, extensions, roots), list)

    def test_index_manager_worker_and_quick_ui_preserve_live_search(self):
        os.environ['QT_QPA_PLATFORM'] = 'offscreen'
        from PyQt6.QtCore import QEventLoop, QTimer
        from PyQt6.QtWidgets import QApplication
        from app.tools.document_search_ui import DocumentSearchMainWindow
        from app.tools.index_manager_ui import IndexManagerDialog
        app = QApplication.instance() or QApplication([])
        source = self.folder / 'insurance.txt'
        source.write_text('보험료 계약', encoding='utf-8')
        dialog = IndexManagerDialog(self.db_path, {'.txt'})
        dialog.start_index(False)
        loop = QEventLoop()
        dialog.worker.finished.connect(loop.quit)
        QTimer.singleShot(5000, loop.quit)
        loop.exec()
        app.processEvents()
        self.assertEqual(self.db.index_stats()['indexed'], 1)
        self.assertIn('보험료', source.read_text(encoding='utf-8'))
        dialog.close()

        window = DocumentSearchMainWindow()
        window.index_db_path = self.db_path
        window.search_input.setText('보험료')
        with patch.object(window.integrated_previewer, 'preview_file') as preview:
            window._start_search()
            loop = QEventLoop()
            window.indexed_search_worker.finished.connect(loop.quit)
            QTimer.singleShot(5000, loop.quit)
            loop.exec()
            app.processEvents()
            self.assertEqual(window.result_model.rowCount(), 1)
            preview.assert_not_called()
            window.result_view.setCurrentIndex(window.result_proxy.index(0, 0))
            preview.assert_called_once()
        window.live_mode_radio.setChecked(True)
        self.assertTrue(window.live_mode_radio.isChecked())
        self.assertFalse(window.fast_mode_radio.isChecked())
        window.checked_folder_paths_set.add(str(self.folder))
        window._start_search()
        loop = QEventLoop()
        window.search_worker.finished.connect(loop.quit)
        QTimer.singleShot(5000, loop.quit)
        loop.exec()
        app.processEvents()
        self.assertEqual(window.result_model.rowCount(), 1)
        window.close()

    def test_hwpx_and_cell_use_existing_extractors(self):
        import zipfile
        hwpx = self.folder / 'local.hwpx'
        cell = self.folder / 'local.cell'
        with zipfile.ZipFile(hwpx, 'w') as archive:
            archive.writestr('Contents/section0.xml', '<root><text>hwpxterm</text></root>')
        with zipfile.ZipFile(cell, 'w') as archive:
            archive.writestr('worksheets/sheet1.xml', '<root><v>cellterm</v></root>')
        counts = DocumentIndexer(str(self.db_path)).update_registered({'.hwpx', '.cell'})
        self.assertEqual(counts['new'], 2)
        for term, ext in (('hwpxterm', '.hwpx'), ('cellterm', '.cell')):
            self.assertEqual(len(self.db.search_indexed(term, {ext}, self.db.list_folders())), 1)

    def test_pdf_docx_xlsx_incremental_and_no_ocr(self):
        import pymupdf
        import docx
        import openpyxl
        pdf_file = self.folder / 'summary.pdf'
        pdf = pymupdf.open()
        pdf.new_page().insert_text((50, 60), 'pdftoken')
        pdf.save(pdf_file)
        pdf.close()
        doc_file = self.folder / 'report.docx'
        word = docx.Document()
        word.add_paragraph('oldword')
        word.save(doc_file)
        sheet_file = self.folder / 'table.xlsx'
        workbook = openpyxl.Workbook()
        workbook.active['A1'] = 'sheetword'
        workbook.save(sheet_file)
        workbook.close()
        image_pdf = self.folder / 'scan.pdf'
        pdf = pymupdf.open()
        pdf.new_page()
        pdf.save(image_pdf)
        pdf.close()
        (self.folder / 'not_supported.bin').write_bytes(b'local only')
        indexer = DocumentIndexer(str(self.db_path))
        first = indexer.update_registered({'.pdf', '.docx', '.xlsx'})
        self.assertEqual((first['new'], first['failed'], first['unsupported']), (3, 1, 1))
        for term, ext in (('pdftoken', '.pdf'), ('oldword', '.docx'), ('sheetword', '.xlsx')):
            self.assertEqual(len(self.db.search_indexed(term, {ext}, self.db.list_folders())), 1)
        self.assertEqual(self.db.index_stats()['failed'], 1)
        word = docx.Document()
        word.add_paragraph('newword')
        word.save(doc_file)
        stat = doc_file.stat()
        os.utime(doc_file, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000_000))
        changed = indexer.update_registered({'.pdf', '.docx', '.xlsx'})
        self.assertEqual(changed['modified'], 1)
        self.assertEqual(changed['unchanged'], 2)
        self.assertEqual(self.db.search_indexed('oldword', {'.docx'}, self.db.list_folders()), [])
        self.assertEqual(len(self.db.search_indexed('newword', {'.docx'}, self.db.list_folders())), 1)

    def test_folder_overlap_removal_and_clear_preserve_source_files(self):
        child = self.folder / 'subfolder'
        child.mkdir()
        self.assertFalse(self.db.add_folder(str(child)))
        source = child / 'nested.txt'
        source.write_text('nested content', encoding='utf-8')
        DocumentIndexer(str(self.db_path)).update_registered({'.txt'})
        self.assertEqual(self.db.index_stats()['indexed'], 1)
        self.db.remove_folder(str(self.folder))
        self.assertTrue(source.is_file())
        self.assertEqual(self.db.index_stats()['total'], 0)
        self.db.add_folder(str(child))
        self.db.add_folder(str(self.folder))
        self.assertEqual(self.db.list_folders(), [str(self.folder)])
        DocumentIndexer(str(self.db_path)).update_registered({'.txt'})
        self.db.clear_index()
        self.assertEqual(self.db.index_stats()['total'], 0)
        self.assertEqual(self.db.list_folders(), [str(self.folder)])
        self.assertTrue(source.is_file())

    def test_cancel_keeps_existing_rows_and_next_update_recovers(self):
        first = self.folder / 'old.txt'
        second = self.folder / 'new.txt'
        first.write_text('old body', encoding='utf-8')
        indexer = DocumentIndexer(str(self.db_path))
        indexer.update_registered({'.txt'})
        first.unlink()
        second.write_text('new body', encoding='utf-8')
        cancelled = [False]

        def progress(stats, path):
            if stats['processed']:
                cancelled[0] = True

        partial = indexer.update_registered({'.txt'}, cancelled=lambda: cancelled[0], progress=progress)
        self.assertTrue(partial['cancelled'])
        with closing(sqlite3.connect(self.db_path)) as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM documents').fetchone()[0], 2)
        completed = indexer.update_registered({'.txt'})
        self.assertEqual(completed['deleted'], 1)
        self.assertEqual(self.db.index_stats()['indexed'], 1)

    def test_old_database_migrates_without_discarding_content(self):
        legacy = self.root / 'legacy.db'
        with closing(sqlite3.connect(legacy)) as conn, conn:
            conn.executescript('''
                CREATE TABLE documents (doc_id INTEGER PRIMARY KEY, file_path TEXT UNIQUE NOT NULL,
                    file_name TEXT, file_ext TEXT, file_size INTEGER, mtime TIMESTAMP,
                    indexed_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP);
                CREATE VIRTUAL TABLE fts_index USING fts5(doc_text, tokenize='unicode61');
                CREATE TABLE doc_fts_map (doc_id INTEGER PRIMARY KEY, fts_rowid INTEGER UNIQUE);
                INSERT INTO documents(doc_id, file_path, file_name, file_ext, file_size)
                    VALUES (1, 'C:/legacy.txt', 'legacy.txt', '.txt', 10);
                INSERT INTO fts_index(rowid, doc_text) VALUES (1, 'archive content');
                INSERT INTO doc_fts_map(doc_id, fts_rowid) VALUES (1, 1);
            ''')
        FTS5Database(str(legacy))
        with closing(sqlite3.connect(legacy)) as conn, conn:
            self.assertEqual(conn.execute('SELECT doc_text FROM indexed_fts WHERE rowid = 1').fetchone()[0],
                             'archive content')
            self.assertIn('mtime_ns', [column[1] for column in conn.execute('PRAGMA table_info(documents)')])
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM fts_index').fetchone()[0], 1)


if __name__ == '__main__':
    unittest.main()
